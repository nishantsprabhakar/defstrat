"""
Exports investor summary reports to .docx (Word) format.
Uses python-docx with custom styles for a professional look.
"""

import io
from datetime import datetime
from typing import Optional

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import pandas as pd


# ── Colour constants (RGB) ────────────────────────────────────────────────────
NAVY = RGBColor(0x1a, 0x23, 0x7e)
DARK_BLUE = RGBColor(0x28, 0x35, 0x93)
ACCENT_BLUE = RGBColor(0x15, 0x65, 0xc0)
GREEN = RGBColor(0x2e, 0x7d, 0x32)
AMBER = RGBColor(0xf5, 0x7f, 0x17)
RED = RGBColor(0xc6, 0x28, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF5, 0xF7, 0xFA)
TEXT = RGBColor(0x21, 0x21, 0x21)


def _set_cell_bg(cell, hex_color: str):
    """Set table cell background colour via XML."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _bold_para(doc, text: str, size=11, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def _divider(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A237E")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def _metric_row(table, label: str, value, color=None):
    row = table.add_row()
    label_cell = row.cells[0]
    value_cell = row.cells[1]
    label_cell.text = label
    label_cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    label_cell.paragraphs[0].runs[0].font.size = Pt(9)

    val_str = str(value) if value is not None else "—"
    value_cell.text = val_str
    run = value_cell.paragraphs[0].runs[0]
    run.bold = True
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = color


def _cr(val) -> str:
    if val is None:
        return "—"
    if val >= 100000:
        return f"₹{val/100000:.1f}L Cr"
    if val >= 1000:
        return f"₹{val/1000:.1f}K Cr"
    return f"₹{val:,.0f} Cr"


def _pct(val) -> str:
    return f"{val:.1f}%" if val is not None else "—"


def _num(val, decimals=2) -> str:
    return f"{val:.{decimals}f}x" if val is not None else "—"


# ── Cover page ────────────────────────────────────────────────────────────────

def _add_cover(doc: Document, title: str, subtitle: str):
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("INDIA DEFENCE SECTOR")
    run.font.size = Pt(9)
    run.font.color.rgb = ACCENT_BLUE
    run.bold = True
    run.font.all_caps = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.size = Pt(26)
    run.font.color.rgb = NAVY
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Generated: {datetime.now().strftime('%d %B %Y')}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"© {datetime.now().year} Nishant Prabhakar")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    run.italic = True
    doc.add_page_break()


# ── Single-company report ─────────────────────────────────────────────────────

def build_company_docx(
    metrics: dict,
    chart_bytes: dict,
    trend_df: pd.DataFrame,
    quarterly_df: pd.DataFrame,
    annual_fins: pd.DataFrame,
) -> bytes:
    """
    Build a single-company investor report docx.

    chart_bytes keys expected:
        revenue_pat_bar, quarterly_trend, margin_waterfall,
        shareholding_donut
    """
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    _add_cover(doc, metrics.get("name", "Company"), "Investor Data Summary Report")

    # ── Overview section ──
    doc.add_heading("Company Overview", level=1)
    p = doc.add_paragraph(metrics.get("description", ""))
    p.runs[0].font.size = Pt(10)

    # Key metrics table
    _divider(doc)
    _bold_para(doc, "Key Metrics at a Glance", size=11, color=NAVY)
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    _metric_row(tbl, "Market Cap", _cr(metrics.get("market_cap_cr")))
    _metric_row(tbl, "Current Price", f"₹{metrics.get('current_price') or '—'}")
    _metric_row(tbl, "52-Week High / Low",
                f"₹{metrics.get('52w_high') or '—'} / ₹{metrics.get('52w_low') or '—'}")
    _metric_row(tbl, "P/E Ratio (TTM)", _num(metrics.get("pe_ratio"), 1))
    _metric_row(tbl, "P/B Ratio", _num(metrics.get("pb_ratio"), 1))
    _metric_row(tbl, "EV/EBITDA", _num(metrics.get("ev_ebitda"), 1))
    _metric_row(tbl, "Beta", _num(metrics.get("beta"), 2))
    _metric_row(tbl, "Dividend Yield", _pct(metrics.get("dividend_yield")))
    _metric_row(tbl, "EPS (TTM)", f"₹{metrics.get('eps_ttm') or '—'}")
    doc.add_paragraph()

    # ── Financials section ──
    doc.add_heading("Financial Highlights", level=1)

    fin_tbl = doc.add_table(rows=0, cols=2)
    fin_tbl.style = "Table Grid"

    _metric_row(fin_tbl, "Revenue (Annual)", _cr(metrics.get("revenue_cr")))
    _metric_row(fin_tbl, "Revenue Growth (YoY)", _pct(metrics.get("revenue_growth_yoy")),
                color=GREEN if (metrics.get("revenue_growth_yoy") or 0) > 0 else RED)
    _metric_row(fin_tbl, "EBITDA", _cr(metrics.get("ebitda_cr")))
    _metric_row(fin_tbl, "EBITDA Margin", _pct(metrics.get("ebitda_margin")))
    _metric_row(fin_tbl, "EBIT", _cr(metrics.get("ebit_cr")))
    _metric_row(fin_tbl, "PAT (Net Profit)", _cr(metrics.get("pat_cr")))
    _metric_row(fin_tbl, "PAT Margin", _pct(metrics.get("pat_margin")))
    _metric_row(fin_tbl, "Total Debt", _cr(metrics.get("total_debt_cr")))
    _metric_row(fin_tbl, "Cash & Equivalents", _cr(metrics.get("cash_cr")))
    _metric_row(fin_tbl, "Total Equity", _cr(metrics.get("equity_cr")))
    _metric_row(fin_tbl, "Debt/Equity", _num(metrics.get("debt_to_equity"), 2))
    _metric_row(fin_tbl, "ROE", _pct(metrics.get("roe")))
    _metric_row(fin_tbl, "ROA", _pct(metrics.get("roa")))
    doc.add_paragraph()

    # ── Charts ──
    doc.add_heading("Charts & Visualisations", level=1)

    for key, caption in [
        ("revenue_pat_bar", "Annual Revenue & PAT Trend"),
        ("quarterly_trend", "Quarterly Performance Trend"),
        ("margin_waterfall", "Profitability Cascade"),
        ("shareholding_donut", "Shareholding Pattern"),
    ]:
        if key in chart_bytes and chart_bytes[key]:
            _bold_para(doc, caption, size=10, color=DARK_BLUE)
            img_io = io.BytesIO(chart_bytes[key])
            doc.add_picture(img_io, width=Inches(5.8))
            doc.add_paragraph()

    # ── Annual trend table ──
    if not trend_df.empty:
        doc.add_heading("Annual Financial Trend", level=1)
        _add_df_table(doc, trend_df)
        doc.add_paragraph()

    # ── Quarterly trend table ──
    if not quarterly_df.empty:
        doc.add_heading("Quarterly Financial Trend (Last 8 Qtrs)", level=1)
        _add_df_table(doc, quarterly_df)

    # Footer note
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(
        f"© {datetime.now().year} Nishant Prabhakar  ·  India Defence Investor Hub  ·  "
        "Data sourced from Yahoo Finance / public exchanges. "
        "Past performance is not indicative of future results. For informational use only."
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    run.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ── Sector-wide report ────────────────────────────────────────────────────────

def build_sector_docx(all_metrics: list[dict], chart_bytes: dict) -> bytes:
    """
    Build a sector-wide comparative report docx.

    chart_bytes keys:
        market_cap_bubble, revenue_comparison, return_heatmap, roe_roa_scatter
    """
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    _add_cover(
        doc,
        "India Listed Defence Sector",
        "Comprehensive Investor Comparison Report",
    )

    doc.add_heading("Sector Overview", level=1)
    p = doc.add_paragraph(
        f"This report covers {len(all_metrics)} listed Indian defence companies across "
        "Aerospace, Electronics, Shipbuilding, Missiles, Drones, and related sub-sectors. "
        "Data as of market close."
    )
    p.runs[0].font.size = Pt(10)

    # Summary stats
    valid_mc = [m["market_cap_cr"] for m in all_metrics if m.get("market_cap_cr")]
    if valid_mc:
        total_mc = sum(valid_mc)
        _bold_para(doc, f"Combined Market Cap: {_cr(total_mc)}", size=12, color=NAVY)
    doc.add_paragraph()

    # ── Comparison charts ──
    doc.add_heading("Sector Charts", level=1)
    for key, caption in [
        ("revenue_comparison", "Revenue & PAT Comparison Across Companies"),
        ("market_cap_bubble", "Valuation vs Profitability (Bubble Chart)"),
        ("return_heatmap", "1-Year & 5-Year Price Return Heatmap"),
        ("roe_roa_scatter", "ROE vs ROA Scatter"),
    ]:
        if key in chart_bytes and chart_bytes[key]:
            _bold_para(doc, caption, size=10, color=DARK_BLUE)
            img_io = io.BytesIO(chart_bytes[key])
            doc.add_picture(img_io, width=Inches(6))
            doc.add_paragraph()

    # ── Comparison table ──
    doc.add_heading("Company Comparison Table", level=1)
    rows = []
    for m in sorted(all_metrics, key=lambda x: x.get("market_cap_cr") or 0, reverse=True):
        rows.append({
            "Company": m["short"],
            "Sub-sector": m.get("sub_sector", ""),
            "Mkt Cap (Cr)": _cr(m.get("market_cap_cr")),
            "Revenue (Cr)": _cr(m.get("revenue_cr")),
            "PAT (Cr)": _cr(m.get("pat_cr")),
            "PAT Margin": _pct(m.get("pat_margin")),
            "P/E": f"{m['pe_ratio']:.1f}x" if m.get("pe_ratio") else "—",
            "ROE": _pct(m.get("roe")),
            "D/E": _num(m.get("debt_to_equity"), 1),
            "1Y Rtn": _pct(m.get("price_return_1y")),
        })
    _add_df_table(doc, pd.DataFrame(rows))

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(
        "Disclaimer: Data sourced from Yahoo Finance / public exchanges. "
        "For informational purposes only. Not investment advice."
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    run.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _add_df_table(doc: Document, df: pd.DataFrame):
    """Add a pandas DataFrame as a styled Word table."""
    if df.empty:
        return
    cols = df.columns.tolist()
    tbl = doc.add_table(rows=1, cols=len(cols))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header row
    hdr = tbl.rows[0]
    for i, col in enumerate(cols):
        cell = hdr.cells[i]
        cell.text = str(col)
        _set_cell_bg(cell, "1A237E")
        run = cell.paragraphs[0].runs[0]
        run.font.color.rgb = WHITE
        run.font.bold = True
        run.font.size = Pt(8)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Data rows
    for row_idx, (_, row) in enumerate(df.iterrows()):
        tr = tbl.add_row()
        bg = "F5F7FA" if row_idx % 2 == 0 else "FFFFFF"
        for i, col in enumerate(cols):
            cell = tr.cells[i]
            cell.text = str(row[col]) if row[col] is not None else "—"
            _set_cell_bg(cell, bg)
            run = cell.paragraphs[0].runs[0]
            run.font.size = Pt(8)
