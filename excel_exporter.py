"""
Exports financial data to .xlsx with rich formatting, multiple sheets,
and embedded charts using openpyxl.
"""

import io
from datetime import datetime
from typing import Optional

import pandas as pd
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


# ── Style constants ───────────────────────────────────────────────────────────
NAVY_FILL = PatternFill("solid", fgColor="1A237E")
DARK_BLUE_FILL = PatternFill("solid", fgColor="283593")
ACCENT_FILL = PatternFill("solid", fgColor="1565C0")
GREEN_FILL = PatternFill("solid", fgColor="2E7D32")
AMBER_FILL = PatternFill("solid", fgColor="F57F17")
LIGHT_FILL = PatternFill("solid", fgColor="F5F7FA")
ALT_FILL = PatternFill("solid", fgColor="E8EAF6")
WHITE_FONT = Font(color="FFFFFF", bold=True, size=10)
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10, name="Calibri")
TITLE_FONT = Font(color="1A237E", bold=True, size=14, name="Calibri")
SUBTITLE_FONT = Font(color="283593", bold=False, size=11, name="Calibri")
BODY_FONT = Font(size=10, name="Calibri")
BOLD_FONT = Font(bold=True, size=10, name="Calibri")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT_ALIGN = Alignment(horizontal="right", vertical="center")
THIN = Side(style="thin", color="BDBDBD")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

INR_CR_FMT = '#,##0.00\\ "Cr"'
PCT_FMT = '0.00"%"'
RATIO_FMT = '0.00"x"'


def _header_row(ws, row_num: int, headers: list[str], fill=NAVY_FILL):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col, value=h)
        cell.fill = fill
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def _data_row(ws, row_num: int, values: list, alt: bool = False):
    fill = ALT_FILL if alt else PatternFill("solid", fgColor="FFFFFF")
    for col, v in enumerate(values, 1):
        cell = ws.cell(row=row_num, column=col, value=v)
        cell.fill = fill
        cell.font = BODY_FONT
        cell.alignment = CENTER if isinstance(v, (int, float)) else LEFT
        cell.border = THIN_BORDER


def _col_widths(ws, min_width=8, max_width=30):
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=min_width)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, min_width), max_width)


def _cr(val) -> Optional[float]:
    return round(val, 2) if val is not None else None


def _section_title(ws, row: int, title: str, col_span: int = 10):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = TITLE_FONT
    cell.fill = PatternFill("solid", fgColor="E8EAF6")
    cell.alignment = LEFT


# ── Sector summary sheet ─────────────────────────────────────────────────────

def _write_sector_summary(wb: Workbook, all_metrics: list[dict]):
    ws = wb.create_sheet("Sector Summary", 0)
    ws.sheet_view.showGridLines = False

    # Title block
    ws.merge_cells("A1:J1")
    ws["A1"] = "INDIA DEFENCE SECTOR — INVESTOR COMPARISON"
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16, name="Calibri")
    ws["A1"].fill = NAVY_FILL
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:J2")
    ws["A2"] = f"© {datetime.now().year} Nishant Prabhakar  ·  Generated: {datetime.now().strftime('%d %B %Y %H:%M')}  ·  Data via Yahoo Finance — for informational use only"
    ws["A2"].font = Font(color="9E9E9E", size=9)
    ws["A2"].alignment = CENTER
    ws.row_dimensions[2].height = 18

    # Separator
    ws.row_dimensions[3].height = 8

    headers = [
        "Company", "Sub-sector", "Mkt Cap (₹ Cr)", "Revenue (₹ Cr)",
        "PAT (₹ Cr)", "PAT Margin %", "P/E (x)", "ROE %", "D/E (x)",
        "1Y Return %",
    ]
    _header_row(ws, 4, headers)
    ws.row_dimensions[4].height = 22

    sorted_m = sorted(all_metrics, key=lambda x: x.get("market_cap_cr") or 0, reverse=True)
    for i, m in enumerate(sorted_m, 5):
        alt = (i % 2 == 0)
        row = [
            m.get("short", ""),
            m.get("sub_sector", ""),
            _cr(m.get("market_cap_cr")),
            _cr(m.get("revenue_cr")),
            _cr(m.get("pat_cr")),
            m.get("pat_margin"),
            m.get("pe_ratio"),
            m.get("roe"),
            m.get("debt_to_equity"),
            m.get("price_return_1y"),
        ]
        _data_row(ws, i, row, alt)
        # Number formats
        for col_idx, fmt in [(3, INR_CR_FMT), (4, INR_CR_FMT), (5, INR_CR_FMT),
                             (6, PCT_FMT), (7, RATIO_FMT), (8, PCT_FMT),
                             (9, RATIO_FMT), (10, PCT_FMT)]:
            ws.cell(row=i, column=col_idx).number_format = fmt

    # Totals row
    total_row = len(sorted_m) + 5
    ws.cell(row=total_row, column=1, value="TOTAL / AVG")
    ws.cell(row=total_row, column=1).font = Font(bold=True, color="1A237E")
    ws.cell(row=total_row, column=1).fill = LIGHT_FILL

    valid_mc = [m["market_cap_cr"] for m in all_metrics if m.get("market_cap_cr")]
    if valid_mc:
        cell = ws.cell(row=total_row, column=3, value=sum(valid_mc))
        cell.number_format = INR_CR_FMT
        cell.font = Font(bold=True)
        cell.fill = LIGHT_FILL

    _col_widths(ws)

    # ── Revenue bar chart ──
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Revenue Comparison (₹ Cr)"
    chart.y_axis.title = "₹ Crores"
    chart.style = 10
    chart.width = 20
    chart.height = 12

    data_start = 5
    data_end = 5 + len(sorted_m) - 1
    rev_ref = Reference(ws, min_col=4, min_row=4, max_row=data_end)
    names_ref = Reference(ws, min_col=1, min_row=data_start, max_row=data_end)
    chart.add_data(rev_ref, titles_from_data=True)
    chart.set_categories(names_ref)
    ws.add_chart(chart, f"A{total_row + 3}")


# ── Individual company financials sheet ──────────────────────────────────────

def _write_company_financials(
    wb: Workbook,
    metrics: dict,
    trend_df: pd.DataFrame,
    quarterly_df: pd.DataFrame,
    annual_fins: pd.DataFrame,
    annual_bs: pd.DataFrame,
    annual_cf: pd.DataFrame,
):
    short = metrics.get("short", "Co")[:20]
    ws = wb.create_sheet(f"{short} – Financials")
    ws.sheet_view.showGridLines = False

    # Header
    ws.merge_cells("A1:G1")
    ws["A1"] = metrics.get("name", "Company")
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    ws["A1"].fill = NAVY_FILL
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 28

    r = 3
    _section_title(ws, r, "KEY METRICS", 7)
    r += 1

    kv_pairs = [
        ("Market Cap", _cr(metrics.get("market_cap_cr")), INR_CR_FMT),
        ("Current Price (₹)", metrics.get("current_price"), None),
        ("P/E Ratio", metrics.get("pe_ratio"), RATIO_FMT),
        ("P/B Ratio", metrics.get("pb_ratio"), RATIO_FMT),
        ("EV/EBITDA", metrics.get("ev_ebitda"), RATIO_FMT),
        ("Revenue", _cr(metrics.get("revenue_cr")), INR_CR_FMT),
        ("Revenue Growth YoY", metrics.get("revenue_growth_yoy"), PCT_FMT),
        ("EBITDA", _cr(metrics.get("ebitda_cr")), INR_CR_FMT),
        ("EBITDA Margin", metrics.get("ebitda_margin"), PCT_FMT),
        ("PAT", _cr(metrics.get("pat_cr")), INR_CR_FMT),
        ("PAT Margin", metrics.get("pat_margin"), PCT_FMT),
        ("Total Debt", _cr(metrics.get("total_debt_cr")), INR_CR_FMT),
        ("Cash", _cr(metrics.get("cash_cr")), INR_CR_FMT),
        ("Debt/Equity", metrics.get("debt_to_equity"), RATIO_FMT),
        ("ROE", metrics.get("roe"), PCT_FMT),
        ("ROA", metrics.get("roa"), PCT_FMT),
        ("Beta", metrics.get("beta"), "0.00"),
        ("1Y Price Return", metrics.get("price_return_1y"), PCT_FMT),
        ("5Y Price Return", metrics.get("price_return_5y"), PCT_FMT),
    ]

    _header_row(ws, r, ["Metric", "Value"], DARK_BLUE_FILL)
    r += 1
    for i, (label, val, fmt) in enumerate(kv_pairs):
        alt = i % 2 == 0
        ws.cell(row=r, column=1, value=label).font = BODY_FONT
        ws.cell(row=r, column=1).fill = ALT_FILL if alt else PatternFill("solid", fgColor="FFFFFF")
        ws.cell(row=r, column=1).border = THIN_BORDER
        c = ws.cell(row=r, column=2, value=val)
        c.font = BODY_FONT
        c.fill = ALT_FILL if alt else PatternFill("solid", fgColor="FFFFFF")
        c.border = THIN_BORDER
        c.alignment = RIGHT_ALIGN
        if fmt and val is not None:
            c.number_format = fmt
        r += 1

    r += 1

    # ── Annual trend ──
    if not trend_df.empty:
        _section_title(ws, r, "ANNUAL FINANCIAL TREND (₹ Cr)", 7)
        r += 1
        headers = list(trend_df.columns)
        _header_row(ws, r, headers, ACCENT_FILL)
        r += 1
        for i, (_, row) in enumerate(trend_df.iterrows()):
            alt = i % 2 == 0
            for col_idx, (col, val) in enumerate(row.items(), 1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.fill = ALT_FILL if alt else PatternFill("solid", fgColor="FFFFFF")
                c.border = THIN_BORDER
                c.font = BODY_FONT
                if col_idx > 1 and val is not None:
                    c.number_format = INR_CR_FMT
                    c.alignment = RIGHT_ALIGN
            r += 1
        r += 1

        # Bar chart for annual trend
        if len(trend_df) > 1:
            chart = BarChart()
            chart.type = "col"
            chart.title = f"{short} — Annual Revenue & PAT"
            chart.y_axis.title = "₹ Crores"
            chart.style = 10
            chart.width = 16
            chart.height = 10
            data_start = r - len(trend_df) - 1
            data_end = r - 2
            rev_ref = Reference(ws, min_col=2, max_col=3, min_row=data_start, max_row=data_end)
            year_ref = Reference(ws, min_col=1, min_row=data_start + 1, max_row=data_end)
            chart.add_data(rev_ref, titles_from_data=True)
            chart.set_categories(year_ref)
            ws.add_chart(chart, f"D{data_start}")

    # ── Quarterly trend ──
    if not quarterly_df.empty:
        _section_title(ws, r, "QUARTERLY FINANCIAL TREND (₹ Cr)", 7)
        r += 1
        headers = list(quarterly_df.columns)
        _header_row(ws, r, headers, ACCENT_FILL)
        r += 1
        for i, (_, row) in enumerate(quarterly_df.iterrows()):
            alt = i % 2 == 0
            for col_idx, (col, val) in enumerate(row.items(), 1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.fill = ALT_FILL if alt else PatternFill("solid", fgColor="FFFFFF")
                c.border = THIN_BORDER
                c.font = BODY_FONT
                if col_idx > 1 and val is not None:
                    c.number_format = INR_CR_FMT
                    c.alignment = RIGHT_ALIGN
            r += 1

    _col_widths(ws)


# ── Main entry points ─────────────────────────────────────────────────────────

def build_sector_excel(all_metrics: list[dict], all_data: dict) -> bytes:
    """
    Build a multi-sheet Excel workbook covering all companies.

    all_data: dict keyed by ticker → raw data dict from data_fetcher
    """
    wb = Workbook()
    wb.remove(wb.active)

    _write_sector_summary(wb, all_metrics)

    for metrics in all_metrics:
        ticker = metrics["ticker"]
        raw = all_data.get(ticker, {})
        from data_fetcher import build_revenue_trend, build_quarterly_trend
        trend = build_revenue_trend(raw)
        quarterly = build_quarterly_trend(raw)
        _write_company_financials(
            wb, metrics, trend, quarterly,
            raw.get("financials", pd.DataFrame()),
            raw.get("balance_sheet", pd.DataFrame()),
            raw.get("cashflow", pd.DataFrame()),
        )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def build_company_excel(
    metrics: dict,
    raw_data: dict,
    trend_df: pd.DataFrame,
    quarterly_df: pd.DataFrame,
) -> bytes:
    """Single-company Excel export."""
    wb = Workbook()
    wb.remove(wb.active)

    _write_sector_summary(wb, [metrics])
    _write_company_financials(
        wb, metrics, trend_df, quarterly_df,
        raw_data.get("financials", pd.DataFrame()),
        raw_data.get("balance_sheet", pd.DataFrame()),
        raw_data.get("cashflow", pd.DataFrame()),
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
